#!/usr/bin/env python3
"""Add pie charts to the `distribution` sheet of the benchmark workbook.

Builds a helper table (embodiment x view totals, parsed from the "A+B" strings)
then attaches native Excel pie charts for:
  1. episodes per model
  2. episodes per acceleration family
  3. episodes per acceleration method
  4. episodes per embodiment x view
  5. episodes per embodiment
PNG previews of the same charts go to assets/distribution_pies/.
"""
import shutil
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "[BENCHMARK] Accelerating Inference of Video World Models.xlsx"
PNG_DIR = ROOT / "assets" / "distribution_pies"

MODEL_COLS = {"DreamGen": 3, "DreamDojo": 4, "Ctrl-World": 5}
GRID_ROWS = range(3, 9)  # embodiment x view rows


def parse_count(value) -> int:
    """'150+150' -> 300, 'N/a' -> 0."""
    if value is None:
        return 0
    text = str(value).strip()
    if not text or text.lower().startswith("n/a"):
        return 0
    return sum(int(p) for p in text.split("+") if p.strip().isdigit())


def add_pie(ws, title, labels_ref, data_ref, anchor):
    chart = PieChart()
    chart.title = title
    chart.height = 8.5
    chart.width = 14
    chart.add_data(data_ref, titles_from_data=False)
    chart.set_categories(labels_ref)
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True
    chart.dataLabels.showVal = True
    ws.add_chart(chart, anchor)


def main():
    backup = XLSX.with_suffix(".xlsx.bak")
    if not backup.exists():
        shutil.copy2(XLSX, backup)

    wb = openpyxl.load_workbook(XLSX)
    ws = wb["distribution"]

    # rerunnable: drop charts / helper table written by a previous run
    ws._charts.clear()
    for row in ws.iter_rows(min_row=30, max_row=45, max_col=2):
        for cell in row:
            cell.value = None

    # ---- helper table: embodiment x view + embodiment totals -----------------
    grid = []
    for row in GRID_ROWS:
        embodiment = ws.cell(row=row, column=1).value
        view = ws.cell(row=row, column=2).value
        total = sum(parse_count(ws.cell(row=row, column=c).value) for c in MODEL_COLS.values())
        grid.append((f"{embodiment} {view}", str(embodiment), total))

    start = 31
    ws.cell(row=start - 1, column=1, value="Bảng phụ cho biểu đồ (tự sinh)").font = Font(bold=True)
    ws.cell(row=start, column=1, value="embodiment x view").font = Font(bold=True)
    ws.cell(row=start, column=2, value="episodes").font = Font(bold=True)
    for i, (label, _, total) in enumerate(grid, 1):
        ws.cell(row=start + i, column=1, value=label)
        ws.cell(row=start + i, column=2, value=total)
    grid_first, grid_last = start + 1, start + len(grid)

    emb_totals = {}
    for _, embodiment, total in grid:
        emb_totals[embodiment] = emb_totals.get(embodiment, 0) + total
    emb_start = grid_last + 2
    ws.cell(row=emb_start, column=1, value="embodiment").font = Font(bold=True)
    ws.cell(row=emb_start, column=2, value="episodes").font = Font(bold=True)
    for i, (embodiment, total) in enumerate(emb_totals.items(), 1):
        ws.cell(row=emb_start + i, column=1, value=embodiment)
        ws.cell(row=emb_start + i, column=2, value=total)
    emb_first, emb_last = emb_start + 1, emb_start + len(emb_totals)

    # ---- pie charts ---------------------------------------------------------
    specs = [
        ("Episodes theo model (chưa apply acceleration)", 1, 12, 14, 2, "H2"),
        ("Episodes theo nhóm acceleration method", 1, 19, 22, 2, "H20"),
        ("Episodes theo từng acceleration method", 4, 19, 27, 5, "P20"),
        ("Episodes theo embodiment x view", 1, grid_first, grid_last, 2, "H38"),
        ("Episodes theo embodiment", 1, emb_first, emb_last, 2, "P38"),
    ]
    for title, label_col, r1, r2, data_col, anchor in specs:
        labels = Reference(ws, min_col=label_col, min_row=r1, max_row=r2)
        data = Reference(ws, min_col=data_col, min_row=r1, max_row=r2)
        add_pie(ws, title, labels, data, anchor)

    wb.save(XLSX)

    # ---- PNG previews ------------------------------------------------------
    PNG_DIR.mkdir(parents=True, exist_ok=True)
    png_specs = [
        ("by_model", "Episodes theo model", [("DreamGen", 1605), ("DreamDojo", 756), ("Ctrl-World", 805)]),
        ("by_family", "Episodes theo nhóm acceleration", [
            ("Dense", 3166), ("Acceleration caching", 8693),
            ("Sparse attention", 12664), ("Token reduction", 4820)]),
        ("by_method", "Episodes theo method", [
            ("itm", 2410), ("sito", 2410), ("worldcache", 3166), ("fastercache", 2361),
            ("dicache", 3166), ("svg2", 3166), ("svg1", 3166),
            ("radialAttention", 3166), ("Pisa", 3166)]),
        ("by_embodiment_view", "Episodes theo embodiment x view",
         [(label, total) for label, _, total in grid]),
        ("by_embodiment", "Episodes theo embodiment", list(emb_totals.items())),
    ]
    for name, title, pairs in png_specs:
        labels = [p[0] for p in pairs]
        values = [p[1] for p in pairs]
        fig, ax = plt.subplots(figsize=(7, 5.5), dpi=150)
        ax.pie(values, labels=labels, autopct=lambda p: f"{p:.1f}%\n({round(p / 100 * sum(values))})",
               startangle=90, counterclock=False, textprops={"fontsize": 8},
               wedgeprops={"linewidth": 1, "edgecolor": "white"})
        ax.set_title(f"{title} (tổng {sum(values)})", fontsize=11)
        ax.axis("equal")
        fig.tight_layout()
        out = PNG_DIR / f"{name}.png"
        fig.savefig(out)
        plt.close(fig)
        print(f"wrote {out}")

    print(f"wrote 5 pie charts into {XLSX.name} (sheet 'distribution'), backup at {backup.name}")


if __name__ == "__main__":
    main()
