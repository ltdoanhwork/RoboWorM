#!/usr/bin/env python3
"""Dump every number behind the two-panel speedup figure, for one embodiment.

Panel (a) = multi-view tracks + RIGIS, panel (b) = single-view tracks + KIVA %, with
speedup bars per method and a 4-metric mean on the right axis.

* **Legend correctness.** The mean emitted here is Visual Integrity, PSNR, Text Alignment,
  FVD -- PSNR and FVD *replace* the Actor Activity and Motion Continuity that the older
  `nr_mean4` scalar used. Both are emitted so a figure can be checked against whichever
  its caption claims; `nr_mean4` is superseded for new figures.

Two things this script refuses to get wrong, because both have already bitten:

* **FVD is a DOWN metric.** It is inverted (denominator / method) before entering the mean,
  so all four components point the same way. Averaging it raw would reward the worst model.
* **FVD carries two different protocols.** Multi-view reads the shared per-camera pool
  (`fvd_perview_pooled.json`, comparable across configurations); single-view reads the
  per-configuration reference (`BENCHMARK_DR_vs_real.xlsx`), because cross-model
  real-vs-real proves humanoid (2.18x floor) and single_arm (7.1-11.6x) are not one
  distribution. Each panel is normalised inside its own protocol, never across.

Ratios are emitted under BOTH baselines. Prefer `mean_dense` (value / the Dense row of the
same track+split): all four components exist and Dense is exactly 1.0. `mean_real` is only
3 components, because PSNR has no real-video level -- PSNR of the real video against itself
is infinite and the workbook stores '---'.

Usage:
  models/dreamgen/.venv/bin/python assets/scripts/dump_speedup_figure_data.py \
      --embodiment bimanual
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / 'assets' / 'scripts'))

from make_nr_circular_figures import (                                   # noqa: E402
    AUDITS, CELL_OF, EMB_CSV, EMB_XL, METHOD_ORDER, METHOD_TO_SPEEDUP, METRIC_KEY,
    TRACK_ORDER, read_audit_tabs, read_speedup_xlsx, read_workbook,
)

MV_TRACKS = ['Dreamgen MV', 'Ctrl-World MV']
SV_TRACKS = ['Dreamgen SV', 'DreamDojo SV']
METHODS = ['Dense'] + METHOD_ORDER

# Figure label -> (workbook track, xlsx sheet label used by the DR exports)
DR_TRACK = {'Dreamgen MV': 'Dreamgen MV', 'Ctrl-World MV': 'CW-4view MV',
            'Dreamgen SV': 'Dreamgen SV', 'DreamDojo SV': 'DreamDojo SV'}
# CELL_OF maps bimanual/humanoid 'Ctrl-World MV' to the ctrlworld4_* cell, i.e. the
# figure's Ctrl-World multi-view row IS the 4-view grid. single_arm has no 4-view run.
DR_TRACK_SINGLE_ARM = dict(DR_TRACK, **{'Ctrl-World MV': 'Ctrl-World MV'})

MEAN_FIG = ['Visual Integrity', 'PSNR', 'Text Alignment', 'FVD']
DOWN = {'FVD'}
BASELINES = ('dense', 'real')


def num(v):
    """Workbook DR sheets store numbers as TEXT ('14.59') and gaps as '---'."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s or s in ('---', '--', '-', 'n/a', 'N/A'):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def read_dr_sheet(xlsx: Path, sheet: str) -> dict:
    """{(EMB, track, method, split): value} + real-video level, from a DR sheet.

    Layout: 4 columns per split block (label, method, value, blank); embodiment header
    rows are upper-case with 'Method' beside them; a 'Real Video' row precedes each
    track's method rows. Values are TEXT, so every read goes through `num`.

    NOTE the PSNR sheet's 'Real Video' row is '---' by construction: PSNR of the real
    video against itself is infinite, so PSNR has NO real-video reference level and
    cannot be ratio-normalised the way the NR metrics are. That is why the caller must
    supply a dense baseline for PSNR.
    """
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, data_only=True)[sheet]
    out, real = {}, {}
    for split, (ca, cb, cc) in {'makovian': (1, 2, 3),
                                'non_makovian': (5, 6, 7)}.items():
        emb = track = None
        for r in range(1, ws.max_row + 1):
            a = ws.cell(r, ca).value
            a = a.strip() if isinstance(a, str) else a
            b = ws.cell(r, cb).value
            b = b.strip() if isinstance(b, str) else b
            if isinstance(a, str) and a in EMB_XL and b == 'Method':
                emb, track = a, None
                continue
            if emb is None:
                continue
            if isinstance(a, str) and a not in ('Markovian', 'Non – Markovian') and a:
                track = a
            v = num(ws.cell(r, cc).value)
            if b == 'Real Video':
                if v is not None:
                    real[(emb, track, split)] = v
                continue
            if track and isinstance(b, str) and v is not None:
                out[(emb, track, b, split)] = v
    return {'val': out, 'real': real}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument('--embodiment', default='bimanual',
                    choices=sorted(EMB_CSV.values()))
    ap.add_argument('--nr-xlsx', type=Path,
                    default=REPO / 'assets/BENCHMARK_NR_main (3).xlsx')
    ap.add_argument('--dr-xlsx', type=Path,
                    default=REPO / 'assets/BENCHMARK_DR_vs_real.xlsx')
    ap.add_argument('--perview', type=Path,
                    default=REPO / 'metrics/reports/fvd_perview_pooled.json')
    ap.add_argument('--out-prefix', default='metrics/reports/speedup_figure_data')
    args = ap.parse_args()

    emb_xl = {v: k for k, v in EMB_CSV.items()}[args.embodiment]
    dr_track = (DR_TRACK_SINGLE_ARM if args.embodiment == 'single_arm' else DR_TRACK)

    book = read_workbook(args.nr_xlsx)
    audit = read_audit_tabs(args.nr_xlsx)
    speed, repairs = read_speedup_xlsx(args.nr_xlsx)
    for r in repairs:
        print('  NOTE  %s' % r)
    psnr = read_dr_sheet(args.dr_xlsx, 'PSNR-main')
    fvd_sv = read_dr_sheet(args.dr_xlsx, 'FVD-main')
    perview = json.loads(args.perview.read_text())

    rows = []
    missing = []
    for panel, tracks in (('a_multiview', MV_TRACKS), ('b_singleview', SV_TRACKS)):
        akey, aunit, ahib = next((k, u, h) for _s, k, _l, u, h, _t in AUDITS
                                 if (k == 'RIGIS') == (panel == 'a_multiview'))
        for track in tracks:
            dtr = dr_track[track]
            for split in ('makovian', 'non_makovian'):
                # Denominators for this (track, split): the Dense row and the real-video
                # level. Collected once so every method in the group divides by the same
                # thing -- the bug class here is a ratio whose denominator silently
                # changes per row.
                den = {'dense': {}, 'real': {}}
                for key in ('Visual Integrity', 'Text Alignment'):
                    den['dense'][key] = book['val'].get(
                        (emb_xl, track, 'Dense', split, key))
                    den['real'][key] = book['real'].get((emb_xl, split, key))
                den['dense']['PSNR'] = psnr['val'].get((emb_xl, dtr, 'Dense', split))
                den['real']['PSNR'] = psnr['real'].get((emb_xl, dtr, split))
                if panel == 'a_multiview':
                    dcell = perview['cells'].get(
                        '|'.join((args.embodiment, dtr, split, 'Dense')))
                    den['dense']['FVD'] = dcell['fvd'] if dcell else None
                    den['real']['FVD'] = dcell.get('floor') if dcell else None
                else:
                    den['dense']['FVD'] = fvd_sv['val'].get(
                        (emb_xl, dtr, 'Dense', split))
                    den['real']['FVD'] = fvd_sv['real'].get((emb_xl, dtr, split))

                for meth in METHODS:
                    raw = {}
                    for key in ('Visual Integrity', 'Text Alignment'):
                        v = book['val'].get((emb_xl, track, meth, split, key))
                        if v is None:
                            missing.append((panel, track, split, meth, key))
                        else:
                            raw[key] = v
                    v = psnr['val'].get((emb_xl, dtr, meth, split))
                    if v is None:
                        missing.append((panel, track, split, meth, 'PSNR'))
                    else:
                        raw['PSNR'] = v
                    if panel == 'a_multiview':
                        cell = perview['cells'].get(
                            '|'.join((args.embodiment, dtr, split, meth)))
                        v = cell['fvd'] if cell else None
                    else:
                        v = fvd_sv['val'].get((emb_xl, dtr, meth, split))
                    if v is None:
                        missing.append((panel, track, split, meth, 'FVD'))
                    else:
                        raw['FVD'] = v

                    # Ratios under BOTH baselines. FVD is inverted (denominator/method)
                    # so every component points UP; averaging it raw would reward the
                    # worst model. PSNR has no real-video level (real-vs-real PSNR is
                    # infinite, the sheet stores '---'), so its 'real' ratio is absent
                    # and the real-baseline mean is over 3 components, not 4.
                    ratio = {b: {} for b in BASELINES}
                    for b in BASELINES:
                        for key in MEAN_FIG:
                            m, d = raw.get(key), den[b].get(key)
                            if m is None or not d:
                                continue
                            ratio[b][key] = (d / m) if key in DOWN else (m / d)

                    cell_id = CELL_OF.get((args.embodiment, track))
                    sp = speed.get((cell_id, METHOD_TO_SPEEDUP.get(meth, meth)))
                    if sp is None and meth == 'Dense':
                        sp = (1.0, 'reference')
                    av = audit['val'].get((emb_xl, track, meth, split, akey))
                    row = dict(
                        panel=panel, embodiment=args.embodiment, track=track,
                        dr_track=dtr, split=split, method=meth, cell=cell_id,
                        speedup=(round(sp[0], 4) if sp else None),
                        speedup_source=(sp[1] if sp else None),
                        audit_metric=akey, audit_value=av,
                        audit_real=audit['real'].get((emb_xl, split, akey)),
                        audit_higher_better=ahib,
                        **{f'raw_{k}': raw.get(k) for k in MEAN_FIG},
                    )
                    for b in BASELINES:
                        comp = [ratio[b][k] for k in MEAN_FIG if k in ratio[b]]
                        row.update({f'ratio_{b}_{k}': (round(ratio[b][k], 4)
                                                       if k in ratio[b] else None)
                                    for k in MEAN_FIG})
                        row[f'mean_{b}'] = (round(sum(comp) / len(comp), 4)
                                            if comp else None)
                        row[f'n_used_{b}'] = len(comp)
                        row[f'partial_{b}'] = len(comp) < len(MEAN_FIG)
                    row['fvd_protocol'] = (
                        'shared per-camera pool (embodiment x camera x regime)'
                        if panel == 'a_multiview' else 'per-configuration reference')
                    rows.append(row)

    out_csv = REPO / (args.out_prefix + f'_{args.embodiment}.csv')
    out_json = REPO / (args.out_prefix + f'_{args.embodiment}.json')
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    with out_csv.open('w', newline='') as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    out_json.write_text(json.dumps(dict(
        embodiment=args.embodiment,
        mean_definition=dict(
            components=MEAN_FIG, down_metrics=sorted(DOWN),
            note='Visual Integrity + PSNR + Text Alignment + FVD. PSNR and FVD '
                 'REPLACE the Actor Activity and Motion Continuity used by the '
                 'older nr_mean4 scalar in make_nr_circular_figures.py, which is '
                 'superseded for new figures. FVD enters inverted '
                 '(denominator/method) so all components point up.',
            baselines=dict(
                dense='value / Dense of the same track+split. RECOMMENDED for this '
                      'figure: all 4 components exist, Dense is exactly 1.0, and it '
                      'answers "what did this method lose vs the dense run it '
                      'replaces" -- which is what a speedup figure compares.',
                real='value / Real Video. Only 3 components: PSNR has NO real-video '
                     'level (real-vs-real PSNR is infinite; the sheet stores "---"), '
                     'so a real-baseline mean is NOT the 4-metric mean the legend '
                     'claims. Multi-view FVD additionally uses the pool split-half '
                     'floor as its stand-in level, which is a noise floor, not a '
                     '"real video" score.')),
        fvd_protocols=dict(
            a_multiview='shared per-camera pool, key (embodiment, camera, regime); '
                        'comparable across configurations',
            b_singleview='per-configuration reference; NOT comparable across '
                         'configurations (humanoid 2.18x, single_arm 7.1-11.6x floor '
                         'on cross-model real-vs-real)'),
        missing=[dict(zip(('panel', 'track', 'split', 'method', 'field'), m))
                 for m in missing],
        rows=rows), indent=2))

    print(f'\n=== {args.embodiment}: figure data (mean vs DENSE baseline) ===')
    hdr = '%-12s %-14s %-13s %-13s %8s %9s %9s %3s'
    print(hdr % ('panel', 'track', 'split', 'method', 'speedup', 'audit',
                 'mean', 'n'))
    for r in rows:
        print(hdr % (r['panel'][:12], r['track'][:14], r['split'][:13],
                     r['method'][:13],
                     '-' if r['speedup'] is None else '%.2fx' % r['speedup'],
                     '-' if r['audit_value'] is None else '%.3f' % r['audit_value'],
                     '-' if r['mean_dense'] is None else '%.4f' % r['mean_dense'],
                     r['n_used_dense']))
    part = [r for r in rows if r['partial_dense'] and r['n_used_dense']]
    print(f'\n  {len(rows)} rows; {len(missing)} missing source fields; '
          f'{len(part)} rows with a PARTIAL dense mean (<4 components)')
    for r in part:
        have = [k for k in MEAN_FIG if r[f'ratio_dense_{k}'] is not None]
        print(f'    {r["track"]:14s} {r["split"]:13s} {r["method"]:14s} '
              f'n={r["n_used_dense"]} using {have}')
    print(f'  wrote {out_csv}')
    print(f'  wrote {out_json}')


if __name__ == '__main__':
    main()
